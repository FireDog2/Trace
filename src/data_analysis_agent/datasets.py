from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import load_workbook

from data_analysis_agent.domain import Artifact, ArtifactType
from data_analysis_agent.repositories import ArtifactRepository
from data_analysis_agent.workspace import WorkspaceService


SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


class DatasetError(RuntimeError):
    """Base error for dataset ingestion and inspection."""


class UnsupportedDatasetError(DatasetError):
    """Raised when a file format is outside the MVP scope."""


class InvalidDatasetError(DatasetError):
    """Raised when a supported file cannot be parsed as tabular data."""


@dataclass(frozen=True, slots=True)
class DatasetColumn:
    name: str
    data_type: str
    nullable: bool


@dataclass(frozen=True, slots=True)
class DatasetInspection:
    sheets: tuple[str, ...]
    columns: tuple[DatasetColumn, ...]
    row_count: int
    sample: tuple[dict[str, Any], ...]
    null_counts: dict[str, int]
    duplicate_row_count: int


class DatasetService:
    def __init__(
        self,
        workspace: WorkspaceService,
        artifacts: ArtifactRepository,
    ) -> None:
        self._workspace = workspace
        self._artifacts = artifacts

    def add_file(self, project_id: str, source: Path) -> Artifact:
        extension = source.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDatasetError(
                f"Unsupported dataset format: {extension or 'no extension'}"
            )
        imported = self._workspace.import_source(project_id, source)
        try:
            self._read_frame(imported)
        except (OSError, ValueError, pl.exceptions.PolarsError) as exc:
            imported.unlink(missing_ok=True)
            raise InvalidDatasetError(f"Could not parse {source.name}") from exc
        return self._artifacts.register(
            project_id=project_id,
            task_id=None,
            artifact_type=ArtifactType.SOURCE_DATASET,
            path=imported,
            title=source.name,
            description=f"Imported source dataset {source.name}",
            producer_execution_id=None,
            input_artifact_ids=(),
        )

    def inspect(self, artifact_id: str, sample_rows: int = 20) -> DatasetInspection:
        artifact = self._artifacts.get(artifact_id)
        if artifact.artifact_type not in {
            ArtifactType.SOURCE_DATASET,
            ArtifactType.DERIVED_DATASET,
        }:
            raise InvalidDatasetError("Artifact is not a dataset")
        frame, sheets = self._read_frame(artifact.path)
        sample_limit = min(max(sample_rows, 0), 100)
        null_counts = {
            name: int(count)
            for name, count in frame.null_count().row(0, named=True).items()
        }
        columns = tuple(
            DatasetColumn(
                name=name,
                data_type=str(data_type),
                nullable=null_counts[name] > 0,
            )
            for name, data_type in frame.schema.items()
        )
        sample = tuple(
            {
                name: _json_scalar(value)
                for name, value in row.items()
            }
            for row in frame.head(sample_limit).to_dicts()
        )
        duplicate_count = int(frame.is_duplicated().sum()) if frame.height else 0
        return DatasetInspection(
            sheets=sheets,
            columns=columns,
            row_count=frame.height,
            sample=sample,
            null_counts=null_counts,
            duplicate_row_count=duplicate_count,
        )

    def _read_frame(self, path: Path) -> tuple[pl.DataFrame, tuple[str, ...]]:
        extension = path.suffix.lower()
        if extension == ".csv":
            return (
                pl.read_csv(path, try_parse_dates=True, infer_schema_length=10_000),
                (),
            )
        if extension == ".xlsx":
            return _read_excel(path)
        raise UnsupportedDatasetError(f"Unsupported dataset format: {extension}")


def _read_excel(path: Path) -> tuple[pl.DataFrame, tuple[str, ...]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise InvalidDatasetError(f"Could not open workbook {path.name}") from exc
    try:
        sheet_names = tuple(workbook.sheetnames)
        if not sheet_names:
            raise InvalidDatasetError("Workbook contains no worksheets")
        sheet = workbook[sheet_names[0]]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, ())
        if not raw_headers:
            return pl.DataFrame(), sheet_names
        headers = _unique_headers(raw_headers)
        values = [tuple(row[: len(headers)]) for row in rows]
        if not values:
            return pl.DataFrame(schema={header: pl.Null for header in headers}), sheet_names
        return (
            pl.DataFrame(values, schema=headers, orient="row", infer_schema_length=None),
            sheet_names,
        )
    finally:
        workbook.close()


def _unique_headers(raw_headers: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for position, value in enumerate(raw_headers, start=1):
        base = str(value).strip() if value is not None else f"column_{position}"
        base = base or f"column_{position}"
        count = counts.get(base, 0) + 1
        counts[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _json_scalar(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value
