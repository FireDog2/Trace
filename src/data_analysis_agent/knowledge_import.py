from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

SUPPORTED_KNOWLEDGE_DOCUMENTS = frozenset({".csv", ".docx", ".md", ".pdf", ".txt", ".xlsx"})
MAX_DOCUMENT_CHARACTERS = 200_000
MAX_IMPORTED_ENTRIES = 100


class KnowledgeDocumentError(ValueError):
    """Raised when an uploaded knowledge document cannot be read safely."""


@dataclass(frozen=True, slots=True)
class KnowledgeCandidate:
    kind: str
    name: str
    definition: str
    formula: str | None = None
    dimensions: tuple[str, ...] = ()
    synonyms: tuple[str, ...] = ()


def extract_knowledge_candidates(filename: str, content: bytes) -> tuple[KnowledgeCandidate, ...]:
    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in SUPPORTED_KNOWLEDGE_DOCUMENTS:
        raise KnowledgeDocumentError("仅支持 PDF、Word、Markdown、TXT、CSV 和 Excel 文档")
    try:
        text = _extract_text(suffix, content)
    except KnowledgeDocumentError:
        raise
    except Exception as exc:
        raise KnowledgeDocumentError("文档内容无法解析") from exc
    text = _clean_text(text)
    if not text:
        raise KnowledgeDocumentError("文档中没有可导入的文字")
    if len(text) > MAX_DOCUMENT_CHARACTERS:
        text = text[:MAX_DOCUMENT_CHARACTERS]

    candidates = _parse_sections(text)
    if not candidates:
        candidates = (
            KnowledgeCandidate(
                kind=_classify(safe_name, text),
                name=Path(safe_name).stem[:160] or "未命名业务文档",
                definition=text[:4_000],
            ),
        )
    return _deduplicate(candidates)[:MAX_IMPORTED_ENTRIES]


def _extract_text(suffix: str, content: bytes) -> str:
    if suffix in {".md", ".txt", ".csv"}:
        return _decode_text(content)
    if suffix == ".docx":
        return _docx_text(content)
    if suffix == ".xlsx":
        return _xlsx_text(content)
    if suffix == ".pdf":
        return _pdf_text(content)
    raise KnowledgeDocumentError("不支持的文档类型")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise KnowledgeDocumentError("文本文件必须使用 UTF-8 或 GB18030 编码")


def _docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            document = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise KnowledgeDocumentError("Word 文档格式无效") from exc
    root = ElementTree.fromstring(document)
    paragraphs: list[str] = []
    for paragraph in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        value = "".join(
            node.text or ""
            for node in paragraph.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
            )
        ).strip()
        if value:
            paragraphs.append(value)
    return "\n".join(paragraphs)


def _xlsx_text(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip() for value in row if value is not None and str(value).strip()
                ]
                if values:
                    lines.append("：".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def _pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared by the package
        raise KnowledgeDocumentError("当前环境缺少 PDF 解析组件") from exc
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise KnowledgeDocumentError("暂不支持加密 PDF") from exc
    if len(reader.pages) > 300:
        raise KnowledgeDocumentError("PDF 超过 300 页，请拆分后上传")
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)


def _clean_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(lines).strip()


def _parse_sections(text: str) -> tuple[KnowledgeCandidate, ...]:
    heading_pattern = re.compile(r"^#{1,6}\s+(.+?)\s*$", re.MULTILINE)
    headings = list(heading_pattern.finditer(text))
    raw_sections: list[tuple[str | None, str]] = []
    if headings:
        for index, heading_match in enumerate(headings):
            end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
            raw_sections.append(
                (
                    heading_match.group(1).strip(),
                    text[heading_match.end() : end].strip(),
                )
            )
    else:
        blocks = re.split(r"\n\s*\n+", text)
        raw_sections = [(None, block.strip()) for block in blocks if block.strip()]

    parsed: list[KnowledgeCandidate] = []
    for section_heading, body in raw_sections:
        candidate = _parse_section(section_heading, body)
        if candidate is not None:
            parsed.append(candidate)
    return tuple(parsed)


_FIELD_PATTERN = re.compile(
    r"^(类型|分类|名称|指标名称|术语名称|规则名称|定义|口径|说明|释义|公式|计算公式|维度|适用维度|同义词|别名)\s*[：:]\s*(.*)$",
    re.IGNORECASE,
)


def _parse_section(heading: str | None, body: str) -> KnowledgeCandidate | None:
    fields: dict[str, str] = {}
    free_lines: list[str] = []
    current_key: str | None = None
    for raw_line in body.splitlines():
        line = re.sub(r"^[-*•]\s*", "", raw_line).strip()
        if not line:
            continue
        match = _FIELD_PATTERN.match(line)
        if match:
            current_key = match.group(1)
            fields[current_key] = match.group(2).strip()
        elif current_key in {"定义", "口径", "说明", "释义"}:
            fields[current_key] = f"{fields[current_key]}\n{line}".strip()
        else:
            free_lines.append(line)

    name = _first(fields, "名称", "指标名称", "术语名称", "规则名称") or (heading or "")
    definition = _first(fields, "定义", "口径", "说明", "释义")
    if (
        not name
        and free_lines
        and len(free_lines[0]) <= 160
        and (definition or len(free_lines) >= 2)
    ):
        name = free_lines.pop(0)
    if not definition:
        definition = "\n".join(free_lines)
    if not name or not definition:
        return None

    explicit_kind = _first(fields, "类型", "分类")
    kind = _normalize_kind(explicit_kind) if explicit_kind else _classify(name, definition)
    formula = _first(fields, "公式", "计算公式") or None
    dimensions = _split_values(_first(fields, "维度", "适用维度"))
    synonyms = _split_values(_first(fields, "同义词", "别名"))
    return KnowledgeCandidate(
        kind=kind,
        name=name[:160].strip(),
        definition=definition[:4_000].strip(),
        formula=formula[:4_000].strip() if formula else None,
        dimensions=dimensions[:30],
        synonyms=synonyms[:30],
    )


def _first(fields: dict[str, str], *keys: str) -> str:
    return next((fields[key] for key in keys if fields.get(key)), "")


def _normalize_kind(value: str) -> str:
    normalized = value.strip().casefold()
    if normalized in {"metric", "指标", "指标口径"}:
        return "metric"
    if normalized in {"rule", "规则", "业务规则"}:
        return "rule"
    return "term"


def _classify(name: str, definition: str) -> str:
    value = f"{name} {definition}".casefold()
    if any(token in value for token in ("规则", "必须", "不得", "应当", "排除", "仅限")):
        return "rule"
    if any(
        token in value
        for token in (
            "指标",
            "公式",
            "计算",
            "利润",
            "收入",
            "成本",
            "销售额",
            "gmv",
            "kpi",
            "率",
            "metric",
            "revenue",
            "cost",
            "profit",
            "margin",
        )
    ):
        return "metric"
    return "term"


def _split_values(value: str) -> tuple[str, ...]:
    return tuple(
        dict.fromkeys(item.strip() for item in re.split(r"[,，、;；]", value) if item.strip())
    )


def _deduplicate(
    candidates: tuple[KnowledgeCandidate, ...],
) -> tuple[KnowledgeCandidate, ...]:
    output: list[KnowledgeCandidate] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = candidate.name.casefold()
        if key not in seen:
            seen.add(key)
            output.append(candidate)
    return tuple(output)
